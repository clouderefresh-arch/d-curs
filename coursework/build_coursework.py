#!/usr/bin/env python3
"""Сборка курсовой работы по требованиям к оформлению (А4, TNR 14, интервал 1,5)."""

import os
import re

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor, Twips

from content import CONTENT, SOURCES

ROOT = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(ROOT, "figures")
OUT = os.path.join(ROOT, "Курсовая_работа_ЕГРН.docx")
os.makedirs(FIG, exist_ok=True)

font_manager.fontManager.addfont("/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf")
font_manager.fontManager.addfont("/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman_Bold.ttf")
plt.rcParams["font.family"] = "Times New Roman"


def set_run_font(run, size=14, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "E7EEF5")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def set_valign(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center")
    tcPr.append(v)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)


def set_repeat_header(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    trPr.append(h)


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(int(w.twips) for w in widths)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for i, child in enumerate(list(grid)):
            if i < len(widths):
                child.set(qn("w:w"), str(int(widths[i].twips)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = widths[i]


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, 14)


def configure_section(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(10)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)
    section.different_first_page_header_footer = True
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    add_page_number(p)
    fp = section.first_page_header.paragraphs[0]
    fp.text = ""


def paragraph(doc, text, *, size=14, bold=False, italic=False, center=False, indent=True,
              before=0, after=0, page_break=False, spacing=1.5, left=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    if left:
        pf.left_indent = Cm(left)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if page_break:
        pf.page_break_before = True
    p.paragraph_format.widow_control = True
    run = p.add_run(text)
    set_run_font(run, size, bold, italic)
    return p


def fill_cell(cell, text, *, bold=False, center=True, size=12, header=False):
    cell.text = ""
    parts = str(text).split("\n")
    for i, part in enumerate(parts):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        pf.line_spacing = 1.0
        pf.first_line_indent = Cm(0)
        run = p.add_run(part)
        set_run_font(run, size, bold)
    set_cell_border(cell)
    set_valign(cell)
    if header:
        shade_header(cell)


TOKEN = re.compile(r"\{\{([A-Z0-9_]+)(?:\|([^}]+))?\}\}")


def collect_keys(obj, found):
    if isinstance(obj, str):
        for key, _page in TOKEN.findall(obj):
            if key not in found:
                found.append(key)
    elif isinstance(obj, (list, tuple)):
        for item in obj:
            collect_keys(item, found)


def replace_keys(obj, numbers, pages):
    if isinstance(obj, str):
        def repl(match):
            key = match.group(1)
            num = numbers[key]
            page = match.group(2) or pages.get(key)
            if page:
                return f"[{num}, с. {page}]"
            return f"[{num}]"
        return TOKEN.sub(repl, obj)
    if isinstance(obj, list):
        return [replace_keys(x, numbers, pages) for x in obj]
    if isinstance(obj, tuple):
        return tuple(replace_keys(x, numbers, pages) for x in obj)
    return obj


def make_figures():
    fig, ax = plt.subplots(figsize=(6.55, 3.55))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 62)
    ax.axis("off")

    def box(x, y, w, h, text, fc, tc="white", fs=8.0):
        ax.add_patch(
            plt.Rectangle((x, y), w, h, facecolor=fc, edgecolor="#1F4E79", linewidth=0.8, zorder=2)
        )
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", color=tc, fontsize=fs, zorder=3)

    box(27, 48, 46, 10, "Единый государственный\nреестр недвижимости", "#1F4E79", fs=9)
    items = [
        (2, 24, "Реестр объектов\nнедвижимости"),
        (35, 24, "Реестр прав,\nограничений\nи обременений"),
        (68, 24, "Реестр границ"),
        (2, 4, "Реестровые дела"),
        (35, 4, "Кадастровые карты"),
        (68, 4, "Книги учёта\nдокументов"),
    ]
    for x, y, text in items:
        box(x, y, 30, 14, text, "#2E75B6", fs=8)
        ax.annotate(
            "",
            xy=(x + 15, y + 14),
            xytext=(50, 48),
            arrowprops=dict(arrowstyle="-", color="#1F4E79", lw=0.7),
            zorder=1,
        )
    path1 = os.path.join(FIG, "fig1_structure.png")
    fig.tight_layout(pad=0.2)
    fig.savefig(path1, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close()

    years = ["2022", "2023", "2024", "2025"]
    vals = [415.5, 653.9, 781.6, 788.8]
    fig, ax = plt.subplots(figsize=(6.55, 3.35))
    bars = ax.bar(years, vals, color="#1F4E79", width=0.62, zorder=3)
    for bar, val in zip(bars, vals):
        label = f"{val:.1f}".replace(".", ",")
        ax.text(bar.get_x() + bar.get_width() / 2, val + 18, label, ha="center", va="bottom", fontsize=10)
    ax.set_ylabel("тыс. заявлений", fontsize=11)
    ax.set_ylim(0, 980)
    ax.tick_params(labelsize=11)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.yaxis.grid(True, linestyle="--", linewidth=0.4, color="#B0B0B0", zorder=0)
    ax.set_axisbelow(True)
    path2 = os.path.join(FIG, "fig2_moscow.png")
    fig.tight_layout()
    fig.savefig(path2, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close()

    rates = [57.4, 19.5, 0.9]
    labels = ["2023/2022", "2024/2023", "2025/2024"]
    fig, ax = plt.subplots(figsize=(6.55, 3.25))
    bars = ax.bar(labels, rates, color=["#1F4E79", "#2E75B6", "#9DC3E6"], width=0.62, zorder=3)
    for bar, val in zip(bars, rates):
        label = f"{val:.1f}".replace(".", ",")
        ax.text(bar.get_x() + bar.get_width() / 2, val + 1.2, label, ha="center", va="bottom", fontsize=10)
    ax.set_ylabel("темп прироста, %", fontsize=11)
    ax.set_ylim(0, 72)
    ax.tick_params(labelsize=11)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.yaxis.grid(True, linestyle="--", linewidth=0.4, color="#B0B0B0", zorder=0)
    ax.set_axisbelow(True)
    path3 = os.path.join(FIG, "fig3_rates.png")
    fig.tight_layout()
    fig.savefig(path3, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close()

    fig, ax = plt.subplots(figsize=(6.55, 4.3))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 78)
    ax.axis("off")
    box(22, 64, 56, 10, "Участник рынка\n(покупатель, банк, нотариус, орган власти)", "#1F4E79", fs=8)
    box(6, 40, 26, 12, "Запрос сведений\nили заявление\nна учёт и регистрацию", "#2E75B6", fs=7.5)
    box(37, 40, 26, 12, "Росреестр\nправовая экспертиза\nи регистрация", "#2E75B6", fs=7.5)
    box(68, 40, 26, 12, "ППК «Роскадастр»\nкадастровый учёт\nи выписки", "#2E75B6", fs=7.5)
    box(22, 16, 56, 12, "ЕГРН и платформа НСПД\nзапись, выписка, открытые пространственные данные", "#1F4E79", fs=8)
    box(22, 2, 56, 9, "Решение о сделке, залоге, налоге, проекте", "#5B9BD5", tc="#1A1A1A", fs=8)
    for x in (19, 50, 81):
        ax.annotate("", xy=(x, 52), xytext=(50, 64), arrowprops=dict(arrowstyle="->", color="#1F4E79", lw=0.8))
        ax.annotate("", xy=(50, 28), xytext=(x, 40), arrowprops=dict(arrowstyle="->", color="#1F4E79", lw=0.8))
    ax.annotate("", xy=(50, 11), xytext=(50, 16), arrowprops=dict(arrowstyle="->", color="#1F4E79", lw=0.8))
    path4 = os.path.join(FIG, "fig4_flow.png")
    fig.tight_layout(pad=0.15)
    fig.savefig(path4, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close()
    return path1, path2, path3, path4


def add_toc(doc, toc_pages):
    paragraph(doc, "ОГЛАВЛЕНИЕ", bold=True, center=True, indent=False, before=0, after=12, page_break=True)
    entries = [
        ("СПИСОК СОКРАЩЕНИЙ", "ABBR", 0),
        ("ВВЕДЕНИЕ", "INTRO", 0),
        ("1 Теоретические основы функционирования Единого государственного реестра недвижимости как информационной системы рынка недвижимости", "CH1", 0),
        ("1.1 Понятие и предпосылки формирования Единого государственного реестра недвижимости", "S11", 1),
        ("1.2 Нормативно-правовое регулирование ведения Единого государственного реестра недвижимости", "S12", 1),
        ("1.3 Роль Единого государственного реестра недвижимости в информационном обеспечении рынка недвижимости", "S13", 1),
        ("2 Анализ практики использования Единого государственного реестра недвижимости по материалам Росреестра и ППК «Роскадастр»", "CH2", 0),
        ("2.1 Состав сведений реестра и их информационное назначение", "S21", 1),
        ("2.2 Порядок и способы предоставления сведений участникам рынка", "S22", 1),
        ("2.3 Оценка результатов использования сведений", "S23", 1),
        ("3 Проблемы и направления совершенствования Единого государственного реестра недвижимости", "CH3", 0),
        ("3.1 Проблемы, снижающие качество информационного обеспечения", "S31", 1),
        ("3.2 Правовые и организационные меры совершенствования", "S32", 1),
        ("3.3 Цифровая трансформация и качество сведений", "S33", 1),
        ("ЗАКЛЮЧЕНИЕ", "CONCL", 0),
        ("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "BIB", 0),
        ("ПРИЛОЖЕНИЕ А Расчёт показателей динамики электронных заявлений", "APP_A", 0),
        ("ПРИЛОЖЕНИЕ Б Схема движения сведений между участником рынка и реестром", "APP_B", 0),
    ]
    for title, key, level in entries:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(1.25 if level else 0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tab_pos = Cm(17.0)
        pf.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(title)
        set_run_font(run, 14, bold=(level == 0 and not title.startswith("1.") and not title.startswith("2.") and not title.startswith("3.") and not title.startswith("ПРИЛОЖ")))
        run2 = p.add_run("\t" + toc_pages.get(key, "0"))
        set_run_font(run2, 14)


def add_title(doc):
    lines_center = [
        ( "Образовательное частное учреждение высшего образования", False, 0),
        ("«МОСКОВСКИЙ ФИНАНСОВО-ЮРИДИЧЕСКИЙ УНИВЕРСИТЕТ МФЮА»", True, 0),
    ]
    for text, bold, before in lines_center:
        paragraph(doc, text, bold=bold, center=True, indent=False, before=before, after=0)
    paragraph(doc, "Кафедра [наименование кафедры]", center=True, indent=False, before=18, after=0)
    paragraph(doc, "КУРСОВАЯ РАБОТА", bold=True, center=True, indent=False, before=28, after=0)
    paragraph(doc, "по дисциплине «Основы ведения единого государственного реестра недвижимости»", center=True, indent=False, before=12, after=0)
    paragraph(doc, "на тему:", center=True, indent=False, before=16, after=0)
    paragraph(
        doc,
        "«Единый государственный реестр недвижимости как информационное обеспечение рынка недвижимости»",
        bold=True,
        center=True,
        indent=False,
        before=6,
        after=0,
    )
    paragraph(doc, "Выполнил(а): студент(ка) [ФИО]", center=False, indent=False, before=36, after=0, left=7.5)
    paragraph(doc, "Группа: [номер группы]", indent=False, before=6, after=0, left=7.5)
    paragraph(doc, "Научный руководитель: [должность, ФИО]", indent=False, before=6, after=0, left=7.5)
    paragraph(doc, "Москва 2026", center=True, indent=False, before=48, after=0)


def add_abbreviations(doc):
    paragraph(doc, "СПИСОК СОКРАЩЕНИЙ", bold=True, center=True, indent=False, before=12, after=8, page_break=True)
    items = [
        "ГК РФ – Гражданский кодекс Российской Федерации",
        "ЕГРН – Единый государственный реестр недвижимости",
        "ЕГРП – Единый государственный реестр прав на недвижимое имущество и сделок с ним",
        "МФЦ – многофункциональный центр предоставления государственных и муниципальных услуг",
        "НК РФ – Налоговый кодекс Российской Федерации",
        "НСПД – Национальная система пространственных данных",
        "ППК – публично-правовая компания",
        "Росреестр – Федеральная служба государственной регистрации, кадастра и картографии",
    ]
    for item in items:
        paragraph(doc, item, indent=False, before=0, after=2)


def render_table(doc, caption, headers, rows, note, widths):
    paragraph(doc, caption, bold=True, indent=False, before=10, after=4, center=False)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.style = "Table Grid"
    for i, head in enumerate(headers):
        fill_cell(table.rows[0].cells[i], head, bold=True, header=True)
    set_repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for r, row in enumerate(rows):
        prevent_row_split(table.rows[r + 1])
        for c, value in enumerate(row):
            fill_cell(table.rows[r + 1].cells[c], value, center=True, header=False)
    set_table_widths(table, widths)
    if note:
        paragraph(doc, note, size=12, italic=True, indent=False, before=2, after=6, spacing=1.0)


def build(toc_pages):
    figs = make_figures()
    fig_map = {"FIG1": figs[0], "FIG2": figs[1], "FIG3": figs[2], "FIG4": figs[3]}

    found = []
    collect_keys(CONTENT, found)
    missing = [key for key in found if key not in SOURCES]
    if missing:
        raise SystemExit(f"Нет библиографических записей: {missing}")
    numbers = {key: i + 1 for i, key in enumerate(found)}
    pages = {key: SOURCES[key][1] for key in found}
    content = replace_keys(CONTENT, numbers, pages)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    configure_section(doc.sections[0])

    add_title(doc)
    add_toc(doc, toc_pages)
    add_abbreviations(doc)

    for block in content:
        kind = block[0]
        if kind == "h":
            level, text = block[1], block[2]
            paragraph(
                doc,
                text,
                bold=True,
                center=True,
                indent=False,
                before=0 if level == 0 else 12,
                after=12,
                page_break=(level == 0),
            )
        elif kind == "p":
            paragraph(doc, block[1])
        elif kind == "center":
            paragraph(doc, block[1], bold=True, center=True, indent=False, before=6, after=10)
        elif kind == "formula":
            paragraph(doc, block[1], center=True, indent=False, before=6, after=2)
        elif kind == "table":
            _, caption, headers, rows, note, widths = block
            render_table(doc, caption, headers, rows, note, widths)
        elif kind == "fig":
            _, key, caption = block
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(fig_map[key], width=Cm(15.8))
            paragraph(doc, caption, center=True, indent=False, before=2, after=8)
        elif kind == "bib":
            paragraph(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", bold=True, center=True, indent=False, before=0, after=12, page_break=True)
            for i, key in enumerate(found, start=1):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.line_spacing = 1.5
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.left_indent = Cm(1.25)
                pf.first_line_indent = Cm(-1.25)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                run = p.add_run(f"{i}. {SOURCES[key][0]}")
                set_run_font(run, 14)
        else:
            raise SystemExit(f"Неизвестный блок {kind}")

    unused = [key for key in SOURCES if key not in numbers]
    if unused:
        raise SystemExit(f"Источники не упомянуты: {unused}")
    doc.save(OUT)
    print(f"sources {len(found)}")
    print(f"saved {OUT}")


if __name__ == "__main__":
    import json
    import sys

    toc_path = os.path.join(ROOT, "toc_pages.json")
    if len(sys.argv) > 1 and sys.argv[1] == "rebuild" and os.path.exists(toc_path):
        with open(toc_path, encoding="utf-8") as fh:
            toc_pages = json.load(fh)
    else:
        toc_pages = {}
    build(toc_pages)
